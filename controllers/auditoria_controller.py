from flask import render_template,request,send_file,session
from controllers.base_controller import BaseController
from database.connection import BancoMysql
from repositories.auditoria_repository import AuditoriaRepository
from services.auditoria_service import AuditoriaService

# Criação de excel
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime
from io import BytesIO

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4,landscape
from reportlab.platypus import (SimpleDocTemplate,Table,TableStyle)

class AuditoriaController(BaseController):
    def __init__(self,app):
        self.rotas = [
            ('/auditoria','auditoria',self.listar,['GET'],'administrador'),
            ('/auditoria/excel','auditoria_excel',self.exportar_excel,['GET'],'administrador'),
            ('/auditoria/pdf','auditoria_pdf',self.exportar_pdf,['GET'],'administrador')
        ]

        super().__init__(app)
        self.db = BancoMysql()
        self.auditoria_repository = AuditoriaRepository(self.db)
        self.auditoria_service = AuditoriaService(self.auditoria_repository)

    def listar(self):
         acao = request.args.get("acao")
         entidade = request.args.get("entidade")
         usuario_id = request.args.get("usuario_id")
         auditorias = self.auditoria_service.listar(acao=acao,entidade=entidade,usuario_id=usuario_id)
         return render_template("html_auditoria/auditoria.html", auditorias=auditorias,acao=acao,entidade=entidade,usuario_id=usuario_id)
    
    def exportar_excel(self):
        acao = request.args.get("acao")
        entidade = request.args.get("entidade")
        usuario_id = request.args.get("usuario_id")
        auditorias = self.auditoria_service.listar_para_relatorio(acao=acao,entidade=entidade,usuario_id=usuario_id)

        arquivo_excel = Workbook()

        planilha = arquivo_excel.active

        planilha.title = "Auditoria"

        planilha["A1"] = "Relatório de Auditoria"

        planilha["A1"].font = Font(
            bold=True,
            size=16
        )

    

        planilha.merge_cells("A1:I1")


        nomes_colunas = [
            "ID",
            "Usuário ID",
            "nome",
            "Ação",
            "Entidade",
            "Entidade ID",
            "Descrição",
            "IP",
            "Data/Hora"
        ]


 

        for numero_coluna, nome_coluna in enumerate(
            nomes_colunas,
            start=1
        ):

            celula = planilha.cell(
                row=3,
                column=numero_coluna
            )

            celula.value = nome_coluna


            celula.font = Font(
                bold=True,
                color="FFFFFF"
            )

    

            celula.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78"
            )

         

            celula.alignment = Alignment(
                horizontal="center"
            )




        for numero_linha, auditoria in enumerate(
            auditorias,
            start=4
        ):

            planilha.cell(
                row=numero_linha,
                column=1,
                value=auditoria["id"]
            )

            planilha.cell(
                row=numero_linha,
                column=2,
                value=auditoria["usuario_id"]
            )

            planilha.cell(
                row=numero_linha,
                column=3,
                value=auditoria["nome"]
            )

            planilha.cell(
                row=numero_linha,
                column=4,
                value=auditoria["acao"]
            )

            planilha.cell(
                row=numero_linha,
                column=5,
                value=auditoria["entidade"]
            )

            planilha.cell(
                row=numero_linha,
                column=6,
                value=auditoria["entidade_id"]
            )

            planilha.cell(
                row=numero_linha,
                column=7,
                value=auditoria["descricao"]
            )

            planilha.cell(
                row=numero_linha,
                column=8,
                value=auditoria["ip"]
            )

            planilha.cell(
                row=numero_linha,
                column=9,
                value=auditoria["data_hora"]
            )



        larguras_colunas = [
            10,  
            12,  
            20,  
            20,  
            20,  
            12,  
            40, 
            18, 
            22   
        ]


        for numero_coluna, largura in enumerate(
            larguras_colunas,
            start=1
        ):

            letra_coluna = get_column_letter(
                numero_coluna
            )

            planilha.column_dimensions[
                letra_coluna
            ].width = largura



        planilha.freeze_panes = "A4"

        arquivo_memoria = BytesIO()

        arquivo_excel.save(arquivo_memoria)
        arquivo_memoria.seek(0)
        data_hora_atual = datetime.now()

        nome_arquivo=(
            f"relatorio_auditoria_"
            f"{data_hora_atual.strftime('%d-%m-%Y_%H-%M-%S')}"
            f".xlsx"
        )

        return send_file(
            arquivo_memoria,
            as_attachment = True,
            download_name=nome_arquivo,
            mimetype=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            )
        )
    
    def valor_ou_vazio(self,valor):
        #usado para não ficar retornando (valor null)
        return "" if valor is None else str(valor)
    
    def exportar_pdf(self):
        acao = request.args.get("acao")
        entidade = request.args.get("entidade")
        usuario_id = request.args.get("usuario_id")
        auditorias = self.auditoria_service.listar_para_relatorio(acao=acao,entidade=entidade,usuario_id=usuario_id)
        arquivo = BytesIO()
        pdf = SimpleDocTemplate(arquivo,pagesize=landscape(A4))
        dados = []
        dados.append(["nome","acao","entidade","entidade id","descricao","ip","data/hora"])

        for auditoria in auditorias:
            dados.append([
                self.valor_ou_vazio(auditoria["nome"]),
                self.valor_ou_vazio(auditoria["acao"]),
                self.valor_ou_vazio(auditoria["entidade"]),
                self.valor_ou_vazio(auditoria["entidade_id"]),
                self.valor_ou_vazio(auditoria["descricao"]),
                self.valor_ou_vazio(auditoria["ip"]),
                auditoria["data_hora"].strftime("%d/%m/%Y %H:%M:%S")
                  if auditoria["data_hora"] is not None else ""
            ])
        
        tabela = Table(dados,colWidths=[100,70,100,50,280,60,80])

        tabela.setStyle(
            TableStyle([
                ("BACKGROUND",(0, 0),(-1, 0),colors.HexColor("#1F4E78")),
                ("TEXTCOLOR",(0, 0),(-1, 0),colors.white),
                ("FONTNAME",(0, 0),(-1, 0),"Helvetica-Bold"),
                ("FONTSIZE",(0, 0),(-1, -1),7),
                ( "GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ALIGN",(0, 0),(-1, 0),"CENTER"),
            ])
        )

        self.auditoria_service.registrar(
            usuario_id=session.get("usuario_id"),
            acao="RELATORIO_PDF",
            entidade="RELATORIO_AUDITORIA",
            entidade_id = usuario_id,
            descricao=(
                f"relatorio de auditorio em pdf gerado por "
                f"{session.get('usuario')}"
            ),
            ip=request.remote_addr
        )

        pdf.build([tabela])
        arquivo.seek(0)
        nome_arquivo = (f"relatorio_auditoria_"f"{datetime.now().strftime('%d-%m-%Y_%H-%M-%S')}"f".pdf")
        return send_file(arquivo,as_attachment=True,download_name=nome_arquivo,mimetype="application/pdf")